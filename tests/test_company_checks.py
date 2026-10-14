from __future__ import annotations

import asyncio
import io
import threading
import subprocess
import sys
import time
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import httpx
import pytest
from openpyxl import Workbook, load_workbook

from parser.company_checks import JobStore, JobWorker, parse_input, valid_inn
from parser.export_company_checks import write_report


def workbook_bytes(rows=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Организации"
    for i in range(8):
        ws.append([f"Служебная строка {i}"])
    ws.append(["ИНН организации", "Наименование организации", "Статус организации"])
    for row in rows or [("7733475417", "ООО Тест", "Стадия ликвидации")]:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = (
                    "s"  # malicious text supplied as a literal, not a formula
                )
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_import_header_duplicate_invalid_and_formula():
    companies, errors, count = parse_input(
        workbook_bytes(
            [
                (7733475417, "=1+1", "Стадия ликвидации"),
                ("7733475417", "Тест", ""),
                ("0000000000", "Неверно", ""),
            ]
        )
    )
    assert count == 3 and len(companies) == 1 and len(errors) == 1
    assert len(companies[0]["origins"]) == 2
    assert companies[0]["input"]["наименование организации"] == "=1+1"
    assert valid_inn("7733475417")
    assert not valid_inn("7733475418")


def test_real_report_dimensions():
    supplied = Path(r"C:\Users\chirt\Downloads\Telegram Desktop\data (1).xlsx")
    if not supplied.is_file():
        pytest.skip("Local user fixture")
    companies, rejected, count = parse_input(supplied.read_bytes())
    assert count == 2595
    assert len(companies) > 2500
    assert companies[0]["inn"] == "4205208688"
    assert len(companies) + len(rejected) <= count


def test_checkpoint_resume_and_owner(tmp_path, monkeypatch):
    monkeypatch.setenv("COMPANY_CHECK_PAUSE", "0")
    store = JobStore(tmp_path)
    job = store.create(workbook_bytes(), "input.xlsx", "user:1")
    p = store.companies(job["id"])[0]
    p["sources"]["egrul"] = {
        "checked_at": "2026-10-14T10:00:00+00:00",
        "error": "",
        "data": {"inn": p["inn"], "ogrn": "1257700417611"},
    }
    p["ogrn"] = "1257700417611"
    store.save_company(job["id"], p)
    calls = []

    def runner(source, company, job_id):
        calls.append(source)
        assert company["ogrn"] == "1257700417611"
        return {"error": "captcha"} if source == "kad" else {"inn": company["inn"]}

    worker = JobWorker(store, tmp_path, runner)
    worker.process(job["id"])
    info = store.get(job["id"], "user:1")
    assert info["status"] == "completed_with_errors"
    assert info["progress"] == 100 and info["failed_checks"] == 1
    assert "egrul" not in calls
    with pytest.raises(KeyError):
        store.get(job["id"], "user:2")
    store.resume(job["id"], "user:1", retry_errors=True)
    calls.clear()
    worker.process(job["id"])
    assert calls == ["kad"]


def test_cancel_preserves_partial_results(tmp_path, monkeypatch):
    monkeypatch.setenv("COMPANY_CHECK_PAUSE", "0")
    store = JobStore(tmp_path)
    job = store.create(workbook_bytes(), "input.xlsx", "owner")

    def runner(source, company, job_id):
        if source == "buh":
            store.cancel(job_id, "owner")
        return {"inn": company["inn"]}

    JobWorker(store, tmp_path, runner).process(job["id"])
    info = store.get(job["id"], "owner")
    assert info["status"] == "cancelled" and info["result_ready"]
    assert info["checks_done"] == 1


def test_typed_report_conflicts_and_missing(tmp_path):
    companies, _, _ = parse_input(
        workbook_bytes([("7733475417", "=HYPERLINK(1)", "Стадия ликвидации")])
    )
    p = companies[0]

    def record(d):
        return {"data": d, "error": "", "checked_at": "2026-10-14T10:00:00+00:00"}

    p["sources"] = {
        "egrul": record({"reg_date": "10.05.2020", "status": "действующая"}),
        "buh": record(
            {
                "years": [
                    {
                        "period": "2025",
                        "revenue": 0,
                        "net_profit": -1500,
                        "payables": None,
                    }
                ]
            }
        ),
        "companium": record({"court_cases": 4, "enforcements": None}),
        "checko": record({"court_cases": 0}),
    }
    target = tmp_path / "result.xlsx"
    write_report(companies, [], target)
    wb = load_workbook(target)
    sheet = wb["Компании"]
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert row["ИНН"] == "7733475417"
    assert row["Выручка 2025, ₽"] == 0
    assert row["Кредиторская задолженность 2025, ₽"] is None
    assert row["Арбитражные дела, количество"] == 4
    assert row["Дела: конфликт"] == "Расхождение источников"
    assert row["Категория"] == "Выявлены риски"  # input liquidation is not erased
    assert row["Название"].startswith("'=HYPERLINK")
    assert sheet.auto_filter.ref and sheet.freeze_panes == "C2"
    assert all(c.data_type != "f" for s in wb for r in s for c in r)
    wb.close()


def test_http_upload_authorization_and_range(tmp_path):
    from parser.serve_export import ExportHandler

    with patch("parser.serve_export.LAVOK_INGEST_TOKEN", "test-token"):
        server = ThreadingHTTPServer(("127.0.0.1", 0), ExportHandler)
        server.job_store = JobStore(tmp_path)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        url = f"http://127.0.0.1:{server.server_port}/lavok/checks"
        auth = {"X-Lavok-Ingest-Token": "test-token", "X-Lavok-Owner": "user:1"}
        try:
            with httpx.Client(trust_env=False) as client:
                assert client.get(url).status_code == 403
                created = client.post(url, headers=auth, content=workbook_bytes())
                assert created.status_code == 202
                job = created.json()
                other = dict(auth, **{"X-Lavok-Owner": "user:2"})
                assert (
                    client.get(f"{url}/{job['id']}", headers=other).status_code == 404
                )
                server.job_store.result_path(job["id"]).write_bytes(b"x" * 3000)
                server.job_store.update(job["id"], status="completed")
                result = client.get(
                    f"{url}/{job['id']}/result.xlsx",
                    headers={**auth, "Range": "bytes=0-1023"},
                )
                assert result.status_code == 206 and len(result.content) == 1024
        finally:
            server.shutdown()
            server.server_close()


def test_zsk_does_not_attach_late_answer_to_another_inn():
    from parser.enrich.zsk_bot import _fetch_one

    class Client:
        reads = 0

        async def send_message(self, entity, inn):
            return SimpleNamespace(id=11)

        async def iter_messages(self, entity, limit):
            self.reads += 1
            if limit == 1:
                yield SimpleNamespace(id=10)
            else:
                yield SimpleNamespace(
                    id=13,
                    out=False,
                    message="9716003016 Текущий уровень риска ЗСК: Высокий",
                    reply_to_msg_id=None,
                )
                yield SimpleNamespace(
                    id=12,
                    out=False,
                    message="7733475417 Текущий уровень риска ЗСК: Низкий",
                    reply_to_msg_id=None,
                )

    report = asyncio.run(
        _fetch_one(Client(), None, "7733475417", require_correlation=True, timeout=3)
    )
    assert report.level == "green" and "9716003016" not in report.raw


def test_source_timeout_kills_child(tmp_path, monkeypatch):
    monkeypatch.setenv("COMPANY_CHECK_SOURCE_TIMEOUT", "0.2")
    store = JobStore(tmp_path)
    job = store.create(workbook_bytes(), "input.xlsx", "owner")
    original = subprocess.Popen
    children = []

    def spawn(_command, **kwargs):
        child = original(
            [sys.executable, "-c", "import time; time.sleep(100)"], **kwargs
        )
        children.append(child)
        return child

    monkeypatch.setattr("parser.company_checks.subprocess.Popen", spawn)
    worker = JobWorker(store, tmp_path)
    result = worker.run_source("egrul", store.companies(job["id"])[0], job["id"])
    assert result["error"] == "source_timeout"
    assert children[0].poll() is not None


def test_finance_preserves_zero_and_missing():
    from parser.enrich.buh import _parse_year_block

    year = _parse_year_block(
        {
            "period": "2025",
            "correction": {
                "financialResult": {"current2110": 0, "current2400": -12},
                "balance": {"current1600": 500, "current1250": 3, "current1520": None},
            },
        }
    )
    assert year.revenue == 0 and year.net_profit == -12000
    assert year.assets == 500000 and year.cash == 3000
    assert year.payables is None


def test_restart_resumes_checkpoint(tmp_path, monkeypatch):
    monkeypatch.setenv("COMPANY_CHECK_PAUSE", "0")
    store = JobStore(tmp_path)
    job = store.create(workbook_bytes(), "input.xlsx", "owner")
    store.update(job["id"], status="running")
    company = store.companies(job["id"])[0]
    company["sources"]["egrul"] = {"data": {}, "error": "", "checked_at": "2026-10-14"}
    store.save_company(job["id"], company)
    calls = []

    def runner(source, company, job_id):
        calls.append(source)
        return {"inn": company["inn"]}

    worker = JobWorker(store, tmp_path, runner)
    worker.start()
    try:
        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            if store.get(job["id"], "owner")["status"] == "completed":
                break
            time.sleep(0.1)
        assert store.get(job["id"], "owner")["status"] == "completed"
        assert "egrul" not in calls
    finally:
        worker.stopping.set()
        worker.thread.join(3)
