"""Durable file jobs. One worker, per-source checkpoints, bounded subprocesses."""

from __future__ import annotations

import io
import json
import math
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from parser.company_check_sources import SOURCES, SOURCE_NAMES
from parser.check_process import contain_process

MAX_UPLOAD = 20 * 1024 * 1024
MAX_ROWS = 10000
ACTIVE = ("queued", "waiting", "running", "exporting")


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def valid_inn(value: str) -> bool:
    if not re.fullmatch(r"\d{10}|\d{12}", value) or set(value) == {"0"}:
        return False
    n = list(map(int, value))
    if len(n) == 10:
        return (
            sum(a * b for a, b in zip(n[:9], (2, 4, 10, 3, 5, 9, 4, 6, 8))) % 11 % 10
            == n[9]
        )
    return (
        sum(a * b for a, b in zip(n[:10], (7, 2, 4, 10, 3, 5, 9, 4, 6, 8))) % 11 % 10
        == n[10]
        and sum(a * b for a, b in zip(n[:11], (3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8)))
        % 11
        % 10
        == n[11]
    )


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def header(value) -> str:
    return re.sub(r"\s+", " ", cell_text(value).lower().replace("ё", "е"))


def parse_input(content: bytes) -> tuple[list[dict], list[dict], int]:
    if not content or len(content) > MAX_UPLOAD:
        raise ValueError("Нужен непустой xlsx размером до 20 МБ")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Распакованный файл превышает 100 МБ")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ValueError("Не удалось прочитать xlsx") from exc
    companies: dict[str, dict] = {}
    rejected: list[dict] = []
    imported_rows = 0
    try:
        for sheet in workbook:
            sheet.reset_dimensions()  # supplied report incorrectly declares A1:A1
            columns = None
            metadata = {}
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_number > MAX_ROWS + 50:
                    raise ValueError("Слишком большой диапазон строк листа")
                if columns is None:
                    if row_number > 50:
                        break
                    names = [header(v) for v in row]
                    index = next(
                        (
                            i
                            for i, n in enumerate(names)
                            if n in {"инн", "инн организации", "инн компании"}
                        ),
                        None,
                    )
                    if index is not None:
                        columns = (index, names)
                    elif len(row) > 1 and cell_text(row[0]):
                        metadata[header(row[0])] = cell_text(row[1])
                    continue
                if not any(cell_text(v) for v in row):
                    continue
                imported_rows += 1
                if imported_rows > MAX_ROWS:
                    raise ValueError("В одном задании допускается до 10 000 строк")
                index, names = columns
                inn = cell_text(row[index]) if index < len(row) else ""
                original = {
                    n or f"столбец {i + 1}": cell_text(v)
                    for i, (n, v) in enumerate(zip(names, row))
                }
                origin = {"sheet": sheet.title, "row": row_number, "fields": original}
                if not valid_inn(inn):
                    rejected.append(
                        {
                            "inn": inn,
                            "reason": "ИНН отсутствует или не прошёл проверку контрольных цифр",
                            **origin,
                        }
                    )
                    continue
                if inn not in companies:
                    companies[inn] = {
                        "inn": inn,
                        "input": original,
                        "origins": [],
                        "sources": {},
                        "input_metadata": dict(metadata),
                    }
                companies[inn]["origins"].append(origin)
        if not companies:
            raise ValueError(
                "Не найдено компаний с корректным ИНН. Нужен столбец ИНН или ИНН организации"
            )
        return list(companies.values()), rejected, imported_rows
    finally:
        workbook.close()


class JobStore:
    def __init__(self, root: Path):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.path = self.root / "jobs.sqlite3"
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS jobs (
                    id TEXT PRIMARY KEY, owner TEXT NOT NULL, filename TEXT NOT NULL,
                    status TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
                    current_inn TEXT, source TEXT, error TEXT, input_rows INTEGER NOT NULL,
                    rejected TEXT NOT NULL, cancel_requested INTEGER NOT NULL DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS companies (
                    job_id TEXT NOT NULL, position INTEGER NOT NULL, inn TEXT NOT NULL,
                    payload TEXT NOT NULL, PRIMARY KEY(job_id, inn)
                );
            """)

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=30)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA journal_mode=WAL")
        try:
            with db:
                yield db
        finally:
            db.close()

    def create(self, content: bytes, filename: str, owner: str) -> dict:
        companies, rejected, rows = parse_input(content)
        job_id = uuid.uuid4().hex
        timestamp = now()
        with self.connect() as db:
            active = db.execute(
                "SELECT COUNT(*) FROM jobs WHERE status IN ('queued','waiting','running','exporting')"
            ).fetchone()[0]
            if active >= 10:
                raise ValueError(
                    "Очередь заполнена: дождитесь завершения текущих проверок"
                )
            db.execute(
                "INSERT INTO jobs(id,owner,filename,status,created_at,updated_at,input_rows,rejected) VALUES(?,?,?,?,?,?,?,?)",
                (
                    job_id,
                    owner,
                    Path(filename.replace("\\", "/")).name[:200],
                    "queued",
                    timestamp,
                    timestamp,
                    rows,
                    json.dumps(rejected, ensure_ascii=False),
                ),
            )
            db.executemany(
                "INSERT INTO companies VALUES(?,?,?,?)",
                [
                    (job_id, i, p["inn"], json.dumps(p, ensure_ascii=False))
                    for i, p in enumerate(companies)
                ],
            )
        return self.get(job_id, owner)

    def companies(self, job_id: str) -> list[dict]:
        with self.connect() as db:
            return [
                json.loads(r[0])
                for r in db.execute(
                    "SELECT payload FROM companies WHERE job_id=? ORDER BY position",
                    (job_id,),
                )
            ]

    def save_company(self, job_id: str, company: dict) -> None:
        with self.connect() as db:
            db.execute(
                "UPDATE companies SET payload=? WHERE job_id=? AND inn=?",
                (json.dumps(company, ensure_ascii=False), job_id, company["inn"]),
            )

    def update(self, job_id: str, **values) -> None:
        allowed = {"status", "current_inn", "source", "error", "cancel_requested"}
        if not set(values) <= allowed:
            raise ValueError("Unknown update field")
        values["updated_at"] = now()
        with self.connect() as db:
            db.execute(
                "UPDATE jobs SET " + ",".join(f"{k}=?" for k in values) + " WHERE id=?",
                (*values.values(), job_id),
            )

    def get(self, job_id: str, owner: str) -> dict:
        with self.connect() as db:
            row = db.execute(
                "SELECT * FROM jobs WHERE id=? AND owner=?", (job_id, owner)
            ).fetchone()
        if row is None:
            raise KeyError(job_id)
        result = dict(row)
        with self.connect() as db:
            # Progress polling must not load every original Excel row and report JSON.
            counts = db.execute(
                """
                SELECT COUNT(*) AS total,
                    COALESCE(SUM((SELECT COUNT(*) FROM json_each(json_extract(c.payload,'$.sources')))),0) AS checked,
                    COALESCE(SUM((SELECT COUNT(*) FROM json_each(json_extract(c.payload,'$.sources')) WHERE
                        COALESCE(json_extract(value,'$.error'),'') != '' OR json_extract(value,'$.partial') = 1)),0) AS errors,
                    COALESCE(SUM((SELECT COUNT(*) FROM json_each(json_extract(c.payload,'$.sources'))) = ?),0) AS completed
                FROM companies c WHERE job_id=?
            """,
                (len(SOURCES), job_id),
            ).fetchone()
        result.update(
            total=counts["total"],
            completed=counts["completed"],
            checks_done=counts["checked"],
            checks_total=counts["total"] * len(SOURCES),
            failed_checks=counts["errors"],
            rejected_rows=len(json.loads(result.pop("rejected"))),
            duplicates=result["input_rows"]
            - counts["total"]
            - len(json.loads(row["rejected"])),
            progress=round(
                counts["checked"] / max(1, counts["total"] * len(SOURCES)) * 100, 1
            ),
            source_label=SOURCE_NAMES.get(result.get("source"), ""),
            result_ready=self.result_path(job_id).is_file()
            and result["status"] not in ACTIVE,
        )
        result.pop("owner")
        result.pop("cancel_requested")
        return result

    def list(self, owner: str) -> list[dict]:
        with self.connect() as db:
            ids = [
                r[0]
                for r in db.execute(
                    "SELECT id FROM jobs WHERE owner=? ORDER BY created_at DESC LIMIT 30",
                    (owner,),
                )
            ]
        return [self.get(i, owner) for i in ids]

    def cancel(self, job_id: str, owner: str) -> dict:
        info = self.get(job_id, owner)
        if info["status"] in ACTIVE:
            self.update(job_id, cancel_requested=1)
        return self.get(job_id, owner)

    def resume(self, job_id: str, owner: str, *, retry_errors=False) -> dict:
        info = self.get(job_id, owner)
        if info["status"] in ACTIVE:
            raise ValueError("Задание уже в очереди или выполняется")
        if retry_errors:
            for p in self.companies(job_id):
                p["sources"] = {
                    k: v
                    for k, v in p["sources"].items()
                    if not v.get("error") and not v.get("partial")
                }
                self.save_company(job_id, p)
        self.update(job_id, status="queued", cancel_requested=0, error=None)
        return self.get(job_id, owner)

    def result_path(self, job_id: str) -> Path:
        if not re.fullmatch(r"[a-f0-9]{32}", job_id):
            raise KeyError(job_id)
        return self.root / f"{job_id}.xlsx"


@contextmanager
def parser_lock(root: Path):
    """Same OS lock as the scheduled parser: never share the Telethon session."""
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes

        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateMutexW.argtypes = (
            wintypes.LPVOID,
            wintypes.BOOL,
            wintypes.LPCWSTR,
        )
        kernel.CreateMutexW.restype = wintypes.HANDLE
        kernel.WaitForSingleObject.argtypes = (wintypes.HANDLE, wintypes.DWORD)
        kernel.ReleaseMutex.argtypes = (wintypes.HANDLE,)
        kernel.CloseHandle.argtypes = (wintypes.HANDLE,)
        handle = kernel.CreateMutexW(None, False, "Global\\FirmParserRunJob")
        if not handle:
            raise OSError("Cannot open parser mutex")
        acquired = kernel.WaitForSingleObject(handle, 0) in (0, 128)
        try:
            yield acquired
        finally:
            if acquired:
                kernel.ReleaseMutex(handle)
            kernel.CloseHandle(handle)
    else:
        import fcntl

        with (root / "company_checks.lock").open("a") as file:
            try:
                fcntl.flock(file, fcntl.LOCK_EX | fcntl.LOCK_NB)
                acquired = True
            except BlockingIOError:
                acquired = False
            try:
                yield acquired
            finally:
                if acquired:
                    fcntl.flock(file, fcntl.LOCK_UN)


class JobWorker:
    def __init__(self, store: JobStore, project_root: Path, runner=None):
        self.store = store
        self.project_root = Path(project_root)
        self.runner = runner or self.run_source
        self.stopping = threading.Event()
        self.thread = None

    def start(self):
        with self.store.connect() as db:
            db.execute(
                "UPDATE jobs SET status='queued' WHERE status IN ('running','waiting','exporting')"
            )
        self.thread = threading.Thread(
            target=self.loop, daemon=True, name="company-checks"
        )
        self.thread.start()

    def cancelled(self, job_id: str) -> bool:
        with self.store.connect() as db:
            return bool(
                db.execute(
                    "SELECT cancel_requested FROM jobs WHERE id=?", (job_id,)
                ).fetchone()[0]
            )

    def run_source(self, source: str, company: dict, job_id: str) -> dict:
        with tempfile.TemporaryDirectory(dir=self.store.root) as directory:
            input_path, output_path = (
                Path(directory) / "input.json",
                Path(directory) / "output.json",
            )
            input_path.write_text(
                json.dumps(company, ensure_ascii=False), encoding="utf-8"
            )
            env = dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8")
            command = [
                sys.executable,
                "-m",
                "parser.company_check_sources",
                source,
                str(input_path),
                str(output_path),
            ]
            flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
            process = subprocess.Popen(
                command,
                cwd=self.project_root,
                env=env,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                creationflags=flags,
                start_new_session=os.name != "nt",
            )
            deadline = time.monotonic() + float(
                os.getenv("COMPANY_CHECK_SOURCE_TIMEOUT", "180")
            )
            with contain_process(process):
                while process.poll() is None:
                    if self.stopping.is_set() or self.cancelled(job_id):
                        return {"error": "cancelled"}
                    if time.monotonic() >= deadline:
                        return {"error": "source_timeout"}
                    self.stopping.wait(0.5)
                if not output_path.is_file():
                    return {"error": f"source_process_exit_{process.returncode}"}
                return json.loads(output_path.read_text(encoding="utf-8"))

    def process(self, job_id: str) -> None:
        self.store.update(job_id, status="running", error=None)
        for p in self.store.companies(job_id):
            for source in SOURCES:
                if self.stopping.is_set():
                    return
                if self.cancelled(job_id):
                    self.export(job_id, "cancelled")
                    return
                if source in p["sources"]:
                    continue
                self.store.update(job_id, current_inn=p["inn"], source=source)
                timestamp = now()
                try:
                    result = self.runner(source, p, job_id)
                except Exception as exc:
                    result = {"error": f"{type(exc).__name__}: {str(exc)[:300]}"}
                if self.stopping.is_set() or self.cancelled(job_id):
                    if not self.stopping.is_set():
                        self.export(job_id, "cancelled")
                    return
                partial = (
                    source == "fedresurs"
                    and bool(result.get("lease_note"))
                    and result.get("lease_hits") is None
                )
                p["sources"][source] = {
                    "checked_at": timestamp,
                    "data": result,
                    "error": result.get("error") or "",
                    "partial": partial,
                }
                if source == "egrul" and not result.get("error"):
                    p["ogrn"] = result.get("ogrn") or ""
                    p["director"] = result.get("director") or ""
                self.store.save_company(job_id, p)
                self.stopping.wait(max(0, float(os.getenv("COMPANY_CHECK_PAUSE", "2"))))
        companies = self.store.companies(job_id)
        errors = any(
            r.get("error") or r.get("partial")
            for p in companies
            for r in p["sources"].values()
        )
        self.export(job_id, "completed_with_errors" if errors else "completed")

    def export(self, job_id: str, final_status: str):
        from parser.export_company_checks import write_report

        self.store.update(job_id, status="exporting")
        with self.store.connect() as db:
            row = db.execute(
                "SELECT rejected FROM jobs WHERE id=?", (job_id,)
            ).fetchone()
        write_report(
            self.store.companies(job_id),
            json.loads(row[0]),
            self.store.result_path(job_id),
        )
        self.store.update(job_id, status=final_status, current_inn=None, source=None)

    def loop(self):
        while not self.stopping.is_set():
            with self.store.connect() as db:
                row = db.execute(
                    "SELECT id FROM jobs WHERE status IN ('queued','waiting') ORDER BY created_at LIMIT 1"
                ).fetchone()
            if row:
                try:
                    if self.cancelled(row[0]):
                        self.export(row[0], "cancelled")
                        continue
                    with parser_lock(self.store.root) as acquired:
                        if acquired:
                            self.process(row[0])
                        else:
                            self.store.update(row[0], status="waiting", source=None)
                except Exception as exc:
                    self.store.update(
                        row[0],
                        status="failed",
                        error=f"{type(exc).__name__}: {str(exc)[:300]}",
                    )
            self.stopping.wait(3)
