"""Sortable reports with typed values, source provenance and explicit coverage."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from parser.company_check_sources import SOURCES, SOURCE_NAMES


def data(p: dict, source: str) -> dict:
    report = p["sources"].get(source) or {}
    return report.get("data") or {}


def usable(p: dict, source: str) -> dict:
    report = p["sources"].get(source) or {}
    return {} if report.get("error") else data(p, source)


def date_value(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(
            tzinfo=None
        )
    except ValueError:
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(str(value), fmt)
            except ValueError:
                pass
    return str(value)


def bool_text(value):
    return "Да" if value is True else "Нет" if value is False else "Нет сведений"


def count_values(p, specs):
    observations = [(SOURCE_NAMES[s], usable(p, s).get(key)) for s, key in specs]
    known = [
        (s, v)
        for s, v in observations
        if isinstance(v, int) and not isinstance(v, bool)
    ]
    if not known:
        return None, "Нет сведений", ""
    values = [v for _, v in known]
    return (
        max(values),
        "; ".join(f"{s}: {v}" for s, v in known),
        "Расхождение источников" if len(set(values)) > 1 else "",
    )


def assessment(p: dict) -> tuple[str, str]:
    original = p["input"]
    eg, comp, check, saby, fed, zsk = (
        usable(p, s)
        for s in ("egrul", "companium", "checko", "saby", "fedresurs", "zsk_bot")
    )
    risks = []
    notes = []
    statuses = [
        original.get("статус организации") or original.get("статус") or "",
        eg.get("status") or "",
    ]
    if any(
        any(m in str(s).lower() for m in ("ликвид", "прекращ", "исключ", "банкрот"))
        for s in statuses
    ):
        risks.append(
            "В исходном файле или ЕГРЮЛ указан статус ликвидации, прекращения либо банкротства"
        )
    if any(s.get("unreliable") is True for s in (comp, check, saby)):
        risks.append("Источник сообщает о недостоверных сведениях")
    if fed.get("is_bankrupt") is True:
        risks.append("В карточке Федресурса есть признак банкротства")
    if zsk.get("level") == "red":
        risks.append("Бот ЗСК сообщил высокий риск")
    elif zsk.get("level") == "yellow":
        notes.append("Бот ЗСК сообщил средний риск")
    courts, _, _ = count_values(
        p,
        (
            ("kad", "cases_found"),
            ("companium", "court_cases"),
            ("checko", "court_cases"),
        ),
    )
    debts, _, _ = count_values(
        p,
        (
            ("fssp", "proceedings"),
            ("companium", "enforcements"),
            ("checko", "enforcements"),
        ),
    )
    if courts:
        notes.append(
            f"Обнаружены арбитражные дела: {courts}; нужны роль компании и содержание дел"
        )
    if debts:
        risks.append(f"Источник сообщает об исполнительных производствах: {debts}")
    errors = [
        SOURCE_NAMES[s]
        for s in SOURCES
        if not p["sources"].get(s)
        or p["sources"][s].get("error")
        or p["sources"][s].get("partial")
    ]
    if errors:
        notes.append("Проверка неполная: " + ", ".join(errors))
    notes.append(
        "Отсутствие признаков в карточке не подтверждает отсутствие банкротства, залогов или долгов во всех реестрах"
    )
    category = (
        "Выявлены риски"
        if risks
        else "Требуется разбор"
        if notes[:-1]
        else "Явные риски не обнаружены"
    )
    return category, ". ".join(risks + notes) + "."


def source_link(p, source):
    d = data(p, source)
    if source == "egrul":
        return "https://egrul.nalog.ru/"
    if source == "buh" and d.get("org_id"):
        return f"https://bo.nalog.gov.ru/organizations-card/{d['org_id']}"
    if source == "zsk_bot":
        return "https://t.me/zskbenefitsarbot"
    for field in ("url", "source"):
        value = str(d.get(field) or "")
        if value.startswith(("https://", "http://")):
            return value
    return {
        "kad": "https://kad.arbitr.ru/",
        "fssp": "https://fssp.gov.ru/iss/ip/",
        "fedresurs": "https://fedresurs.ru/",
    }.get(source, "")


def main_row(p, years):
    original = p["input"]
    eg, comp, check, fed, zsk = (
        usable(p, s) for s in ("egrul", "companium", "checko", "fedresurs", "zsk_bot")
    )
    courts, courts_sources, courts_conflict = count_values(
        p,
        (
            ("kad", "cases_found"),
            ("companium", "court_cases"),
            ("checko", "court_cases"),
        ),
    )
    debts, debts_sources, debts_conflict = count_values(
        p,
        (
            ("fssp", "proceedings"),
            ("companium", "enforcements"),
            ("checko", "enforcements"),
        ),
    )
    flags = [
        (SOURCE_NAMES[s], usable(p, s).get("unreliable"))
        for s in ("companium", "checko", "saby")
    ]
    unreliable = (
        "Да"
        if any(v is True for _, v in flags)
        else "Не отмечена источником"
        if any(v is False for _, v in flags)
        else "Нет сведений"
    )
    category, summary = assessment(p)
    values = {
        "ИНН": p["inn"],
        "ОГРН": eg.get("ogrn"),
        "КПП": eg.get("kpp"),
        "Название": eg.get("name_full")
        or eg.get("name")
        or original.get("наименование организации")
        or original.get("название"),
        "Название в файле": original.get("наименование организации")
        or original.get("название"),
        "Статус в файле": original.get("статус организации") or original.get("статус"),
        "Дата исходного отчёта": date_value(
            (p.get("input_metadata") or {}).get("дата формирования отчета")
        ),
        "Статус ЕГРЮЛ": eg.get("status") or "Нет сведений",
        "Дата регистрации": date_value(eg.get("reg_date")),
        "Адрес": eg.get("address"),
        "Руководитель": eg.get("director"),
        "Регион в файле": original.get("регион нахождения организации"),
        "ОКВЭД ЕГРЮЛ": eg.get("okved"),
        "ОКВЭД заявительный": original.get("оквэд заявительного типа"),
        "ОКВЭД отчётный": original.get("оквэд отчетного типа"),
        "Недостоверность в файле": original.get(
            "признак наличия сведений о недостоверности"
        ),
        "Недостоверность по источникам": unreliable,
        "Основание недостоверности": "; ".join(
            f"{s}: {bool_text(v)}" for s, v in flags
        ),
        "Сотрудники": comp.get("employees"),
        "МСП": comp.get("msp"),
        "Уставный капитал, ₽": comp.get("capital_rub"),
        "Учредители (текст источника)": comp.get("founder"),
        "Налоги (сумма источника), ₽": comp.get("taxes_rub"),
        "Страховые взносы, ₽": comp.get("insurance_rub"),
        "Выручка Companium (текст)": comp.get("revenue_note"),
        "Прибыль Companium (текст)": comp.get("profit_note"),
        "Отчётные периоды БФО": ", ".join(
            str(y.get("period"))
            for y in usable(p, "buh").get("years", [])
            if y.get("period")
        ),
        "Арбитражные дела, количество": courts,
        "Дела: сведения источников": courts_sources,
        "Дела: конфликт": courts_conflict,
        "Исполнительные производства, количество": debts,
        "Производства: сведения источников": debts_sources,
        "Производства: конфликт": debts_conflict,
        "Задолженность ФССП, ₽": usable(p, "fssp").get("debt_total_rub"),
        "Суммы исков, ₽": usable(p, "kad").get("claims_total_rub"),
        "Банкротство: статус Федресурса": fed.get("status"),
        "Банкротство: признак в карточке": bool_text(fed.get("is_bankrupt")),
        "Банкротство: текст Companium": comp.get("bankruptcy_reg"),
        "Публикации Федресурса: количество Companium": comp.get("fedresurs_msgs"),
        "Лизинг / залоги: найденные упоминания": fed.get("lease_hits"),
        "Публикации: результат проверки": fed.get("lease_note"),
        "Дисквалификация: результат": bool_text(
            (fed.get("disqualified") or {}).get("found")
        ),
        "Дисквалификация: примечание": (fed.get("disqualified") or {}).get("note"),
        "Лицензии": comp.get("licenses"),
        "Проверки": comp.get("checks"),
        "Санкции": comp.get("sanctions"),
        "ЗСК": zsk.get("label") or "Нет сведений",
        "ЗСК: ответ бота": zsk.get("raw"),
        "ЗСК: дата проверки": date_value(
            (p["sources"].get("zsk_bot") or {}).get("checked_at")
        ),
        "Категория": category,
        "Вывод": summary,
        "Выполнено проверок": len(p["sources"]),
        "Ошибок источников": sum(
            bool(r.get("error") or r.get("partial")) for r in p["sources"].values()
        ),
        "Повторов в файле": len(p["origins"]) - 1,
    }
    by_year = {str(y.get("period")): y for y in usable(p, "buh").get("years", [])}
    for year in years:
        y = by_year.get(year, {})
        for label, key in (
            ("Выручка", "revenue"),
            ("Прибыль", "net_profit"),
            ("Краткосрочные займы", "borrowed_short"),
            ("Долгосрочные займы", "borrowed_long"),
            ("Кредиторская задолженность", "payables"),
            ("Активы", "assets"),
            ("Денежные средства", "cash"),
        ):
            values[f"{label} {year}, ₽"] = y.get(key)
    for source in SOURCES:
        r = p["sources"].get(source) or {}
        values[f"{SOURCE_NAMES[source]}: результат"] = (
            "Не выполнена"
            if not r
            else "Ошибка"
            if r.get("error")
            else "Частичный ответ"
            if r.get("partial")
            else "Получен ответ"
        )
        values[f"{SOURCE_NAMES[source]}: ошибка"] = r.get("error")
        values[f"{SOURCE_NAMES[source]}: дата"] = date_value(r.get("checked_at"))
        values[f"{SOURCE_NAMES[source]}: ссылка"] = source_link(p, source)
    return values


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        # Excel input never becomes an executable formula.
        if value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value[:32767]
    return value


def flatten(value, prefix=""):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from flatten(child, f"{prefix}.{key}" if prefix else str(key))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from flatten(child, f"{prefix}[{index}]")
    elif isinstance(value, str) and len(value) > 32000:
        for index in range(0, len(value), 32000):
            yield f"{prefix}: часть {index // 32000 + 1}", value[index : index + 32000]
    else:
        yield prefix, value


def make_sheet(wb, name, headers, rows):
    sheet = wb.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([safe_cell(v) for v in row])
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 42
    for c in sheet[1]:
        c.fill = PatternFill("solid", fgColor="183153")
        c.font = Font(color="FFFFFF", bold=True, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[c.column_letter].width = min(
            45, max(16, len(str(c.value)) + 2)
        )
    for row in sheet.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)
            if isinstance(c.value, datetime):
                c.number_format = "dd.mm.yyyy hh:mm"
            elif isinstance(c.value, (int, float)):
                c.number_format = "#,##0.##"
            else:
                c.number_format = "@"
            if isinstance(c.value, str) and c.value.startswith(("https://", "http://")):
                c.hyperlink = c.value
                c.font = Font(color="1769AA", underline="single")
    return sheet


def write_report(companies: list[dict], rejected: list[dict], path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    years = sorted(
        {
            str(y["period"])
            for p in companies
            for y in usable(p, "buh").get("years", [])
            if y.get("period")
        },
        reverse=True,
    )
    main = [main_row(p, years) for p in companies]
    headers = list(main[0]) if main else ["ИНН"]
    sheet = make_sheet(
        wb, "Компании", headers, [[r.get(h) for h in headers] for r in main]
    )
    for i, r in enumerate(main, 2):
        color = (
            "FCE3E3"
            if r["Категория"] == "Выявлены риски"
            else "FFF3CD"
            if r["Категория"] == "Требуется разбор"
            else "E4F2E8"
        )
        for cell in sheet[i]:
            cell.fill = PatternFill("solid", fgColor=color)
    finance = []
    sources = []
    details = []
    originals = []
    cases = []
    publications = []
    for p in companies:
        kd = usable(p, "kad")
        for case in kd.get("cases", []):
            cases.append(
                [
                    p["inn"],
                    case.get("CaseNumber") or case.get("caseNumber"),
                    case.get("CaseId") or case.get("id"),
                    bool_text(kd.get("details_complete")),
                    json.dumps(case, ensure_ascii=False),
                    source_link(p, "kad"),
                ]
            )
        fd = usable(p, "fedresurs").get("extras") or {}
        for publication in fd.get("publications", []):
            if isinstance(publication, dict):
                publications.append(
                    [
                        p["inn"],
                        publication.get("number") or publication.get("id"),
                        publication.get("type") or publication.get("typeName"),
                        publication.get("date") or publication.get("publishDate"),
                        bool_text(fd.get("publications_complete")),
                        json.dumps(publication, ensure_ascii=False),
                        source_link(p, "fedresurs"),
                    ]
                )
        for y in usable(p, "buh").get("years", []):
            finance.append(
                [
                    p["inn"],
                    cell_text_period(y.get("period")),
                    y.get("revenue"),
                    y.get("net_profit"),
                    y.get("borrowed_short"),
                    y.get("borrowed_long"),
                    y.get("payables"),
                    y.get("assets"),
                    y.get("cash"),
                    source_link(p, "buh"),
                ]
            )
        for s in SOURCES:
            r = p["sources"].get(s) or {}
            sources.append(
                [
                    p["inn"],
                    SOURCE_NAMES[s],
                    date_value(r.get("checked_at")),
                    "Ошибка"
                    if r.get("error")
                    else "Частичный ответ"
                    if r.get("partial")
                    else "Получен ответ"
                    if r
                    else "Не выполнена",
                    r.get("error"),
                    source_link(p, s),
                    json.dumps(r.get("data") or {}, ensure_ascii=False),
                ]
            )
            for field, value in flatten(r.get("data") or {}):
                details.append(
                    [
                        p["inn"],
                        SOURCE_NAMES[s],
                        field,
                        value,
                        date_value(r.get("checked_at")),
                    ]
                )
        for origin in p["origins"]:
            for field, value in origin["fields"].items():
                originals.append(
                    [p["inn"], origin["sheet"], origin["row"], field, value]
                )
    make_sheet(
        wb,
        "Финансы по годам",
        [
            "ИНН",
            "Год",
            "Выручка, ₽",
            "Прибыль, ₽",
            "Краткосрочные займы, ₽",
            "Долгосрочные займы, ₽",
            "Кредиторская задолженность, ₽",
            "Активы, ₽",
            "Деньги, ₽",
            "Источник",
        ],
        finance,
    )
    make_sheet(
        wb,
        "Проверки источников",
        ["ИНН", "Источник", "Дата", "Результат", "Ошибка", "Ссылка", "Ответ JSON"],
        sources,
    )
    make_sheet(
        wb,
        "Подробные данные",
        ["ИНН", "Источник", "Поле", "Значение", "Дата проверки"],
        details,
    )
    make_sheet(
        wb,
        "Арбитражные дела",
        [
            "ИНН",
            "Номер дела",
            "ID дела",
            "Получен полный список",
            "Данные дела",
            "Источник",
        ],
        cases,
    )
    make_sheet(
        wb,
        "Публикации",
        [
            "ИНН",
            "Номер публикации",
            "Тип",
            "Дата",
            "Получен полный список",
            "Данные публикации",
            "Источник",
        ],
        publications,
    )
    make_sheet(
        wb, "Исходные строки", ["ИНН", "Лист", "Строка", "Поле", "Значение"], originals
    )
    make_sheet(
        wb,
        "Ошибки импорта",
        ["ИНН", "Лист", "Строка", "Причина", "Исходные данные"],
        [
            [
                r["inn"],
                r["sheet"],
                r["row"],
                r["reason"],
                json.dumps(r["fields"], ensure_ascii=False),
            ]
            for r in rejected
        ],
    )
    make_sheet(
        wb,
        "Как читать",
        ["Поле", "Значение"],
        [
            [
                "Пустая сумма",
                "Сумма не получена. Ноль означает полученное числовое значение 0.",
            ],
            [
                "Даты проверок",
                "Время проверки источников указано в UTC. Даты исходного отчёта и регистрации сохранены как в источнике.",
            ],
            [
                "Количество при конфликте",
                "Показан максимум. Значения каждого источника и расхождения сохранены рядом.",
            ],
            [
                "Получен ответ",
                "Источник ответил. Это не означает, что заполнены все поля или исключены все риски.",
            ],
            [
                "Банкротство",
                "Статус карточки и текст агрегатора не заменяют полную проверку сообщений ЕФРСБ.",
            ],
            [
                "Публикации",
                "Счётчик публикаций не означает число залогов или лизингов. Типы требуют содержимого публикаций.",
            ],
            [
                "Финансы",
                "Суммы БФО приведены в рублях по единицам формы; текст Companium сохранён отдельно.",
            ],
            [
                "ЗСК",
                "Результат Telegram-бота, с проверкой связи ответа и отправленного ИНН.",
            ],
            [
                "Детализация",
                "Сохранены все поля, которые вернули модули источников. Не полученные суммы и детали не выдумываются.",
            ],
        ],
    )
    temp = Path(path).with_suffix(".tmp.xlsx")
    wb.save(temp)
    wb.close()
    temp.replace(path)


def cell_text_period(value):
    text = str(value or "")
    return int(text) if text.isdigit() else text
